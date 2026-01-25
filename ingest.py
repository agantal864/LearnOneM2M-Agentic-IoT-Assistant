# 1. This file scraps all data from the ACME documentation (saves the files to rawData/implementations/acme)
# 2. Convert oneM2M standards TS files to markdown (saves the files to rawData/specifications)
# 3. Scraps important codes in the ACME codebase (saves the files to rawData/implementations/acmeCode)

import os
import re
import time
import requests
from bs4 import BeautifulSoup
from markdownify import markdownify as md
from docx import Document

def cleanMarkdown(text):
    # Remove excessive vertical whitespace (3 or more to 2)
    text = re.sub(r'\n{3,}', '\n\n', text)
    # Remove image placeholders like ![](...) to skip images (image description will be added manually)
    text = re.sub(r'!\[.*?\]\(.*?\)', '', text)
    return text.strip()

def scrap(acmePath, acmeUrls):
    # Create file directories for each nav links
    for navitems, urls in acmeUrls.items():
        category = os.path.join(acmePath, navitems)
        os.makedirs(category, exist_ok=True)
        # Scrap each url and load to its specific file directory 
        for url in urls:
            response = requests.get(url)
            if response.status_code == 200:
                soup = BeautifulSoup(response.text, 'html.parser')
                # Target the main article
                article = soup.find('article')
                if article:
                    # Convert HTML article to Markdown
                    # removes unwanted tags, 'heading_style' makes it clean
                    markdownText = md(str(article), heading_style="ATX", strip=['img', 'script', 'style'])
                    # Clean markdown text            
                    finalText = cleanMarkdown(markdownText)
                    # Save as markdown
                    filename = url.rstrip("/").split("/")[-1] + ".md"
                    # If it's the root index, name it index.md
                    if not filename or filename == "acmecse.net.md":
                        filename = "index.md"
                    
                    with open(os.path.join(category, filename), "w", encoding="utf-8") as f:
                        f.write(f"Source: {url}\n\n")
                        f.write(finalText)
                    print(f"Successfully saved Markdown: {filename}")
                else:
                    print(f" No article content found for {url}")
            else:
                print(f" Failed to reach {url}")
            # Added request interval for server politeness
            time.sleep(2)

def convertSpecs(specPath):
    # Check if path exists
    if not os.path.exists(specPath):
        print(f"Error: Path {specPath} not found.")
        return
    # Process all docx files in the directory
    for filename in os.listdir(specPath):
        if filename.endswith(".docx"):
            print(f"Converting {filename}...")
            fullPath = os.path.join(specPath, filename)            
            try:
                doc = Document(fullPath)
                mdLines = []               
                for para in doc.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue                       
                    # Basic Heading Detection for ETSI/oneM2M docx styles
                    if para.style.name.startswith('Heading 1'):
                        mdLines.append(f"\n# {text}\n")
                    elif para.style.name.startswith('Heading 2'):
                        mdLines.append(f"\n## {text}\n")
                    elif para.style.name.startswith('Heading 3'):
                        mdLines.append(f"\n### {text}\n")
                    else:
                        # Regular paragraph
                        mdLines.append(text)
                finalText = "\n\n".join(mdLines)               
                # Save with .md extension
                newFileName = filename.replace(".docx", ".md")
                with open(os.path.join(specPath, newFileName), "w", encoding="utf-8") as f:
                    f.write(f"Source Document: {filename}\n\n")
                    f.write(finalText)
                print(f"Created: {newFileName}")                
            except Exception as e:
                print(f"Failed to convert {filename}: {e}")

def addRepo(repoPath, acmeRepoFiles):
    for category, urls in acmeRepoFiles.items():
        categoryPath = os.path.join(repoPath, category)
        os.makedirs(categoryPath, exist_ok=True)
        for url in urls:
            # Fetch raw text
            response = requests.get(url)
            if response.status_code == 200:
                filename = url.split("/")[-1]
                with open(os.path.join(categoryPath, filename), "w", encoding="utf-8") as f:
                    f.write(f"Source Code: {url}\n\n")
                    f.write(response.text)
                print(f"Saved Code: {category}/{filename}")
            time.sleep(1) 

if __name__ == "__main__":  
    # Folder paths relative to LearnOneM2M root (acme md files will be saved here)
    acmePath = "rawData/implementations/acme"
    # (converted ts files to md will be saved here)
    specPath = "rawData/specifications"
    repoPath = "rawData/implementations/acmeCode" 
    # Note: acmeUrls could be automated using a web scraper like Selenium,
    # but I've chosen a manual approach to minimize server load and respect ACME's resources.
    acmeUrls = {
    "Home": [
        "https://acmecse.net/", 
        "https://acmecse.net/home/ACME-CSE-introduction/",
        "https://acmecse.net/home/oneM2M-introduction/",
        "https://acmecse.net/home/Supported/",
        "https://acmecse.net/home/SupportedRuntime/",
        "https://acmecse.net/home/Roadmap/",
        "https://acmecse.net/home/License/",
        "https://acmecse.net/home/Contact/"
    ],
    "SetupAndRunning": [
        "https://acmecse.net/setup/Installation/",
        "https://acmecse.net/setup/Running/",
        "https://acmecse.net/setup/Database/",
        "https://acmecse.net/setup/Certificates/",
        "https://acmecse.net/setup/Configuration-introduction/",
        "https://acmecse.net/setup/Configuration-basic/",
        "https://acmecse.net/setup/Configuration-cse/",
        "https://acmecse.net/setup/Configuration-database/",
        "https://acmecse.net/setup/Configuration-logging/",
        "https://acmecse.net/setup/Configuration-coap/",
        "https://acmecse.net/setup/Configuration-http/",
        "https://acmecse.net/setup/Configuration-mqtt/",
        "https://acmecse.net/setup/Configuration-ws/",
        "https://acmecse.net/setup/Configuration-scripting/",
        "https://acmecse.net/setup/Configuration-uis/",
        "https://acmecse.net/setup/Configuration-resources/",
        "https://acmecse.net/setup/Operation-diagrams/",
        "https://acmecse.net/setup/Operation-management/",
        "https://acmecse.net/setup/Operation-mqtt/",
        "https://acmecse.net/setup/Operation-uppertester/",
        "https://acmecse.net/setup/Console/",
        "https://acmecse.net/setup/TextUI/",
        "https://acmecse.net/setup/WebUI/"
    ],
    "Development": [
        "https://acmecse.net/development/Overview/",
        "https://acmecse.net/development/UnitTests/",
        "https://acmecse.net/development/StartupResources/",
        "https://acmecse.net/development/AttributePolicies/",
        "https://acmecse.net/development/FlexContainerPolicies/",
        "https://acmecse.net/development/HelpDocumentation/",
        "https://acmecse.net/development/tools/NotificationServer/",
        "https://acmecse.net/development/tools/OnboardingTool/",
        "https://acmecse.net/development/tools/HashCredentials/",
        "https://acmecse.net/development/tools/ZookeeperTool/",
        "https://acmecse.net/development/ACMEScript/",
        "https://acmecse.net/development/ACMEScript-loading/",
        "https://acmecse.net/development/ACMEScript-operations/",
        "https://acmecse.net/development/ACMEScript-functions/",
        "https://acmecse.net/development/ACMEScript-variables/",
        "https://acmecse.net/development/ACMEScript-metatags/",
        "https://acmecse.net/development/ACMEScript-uppertester/",
        "https://acmecse.net/development/Plugins/",
        "https://acmecse.net/development/Embedding_ACME/",
        "https://acmecse.net/development/DebugMode/",
        "https://acmecse.net/development/TypeChecking/",
        "https://acmecse.net/development/ThirdPartyLibraries/"
    ],
    "HowTos": [
        "https://acmecse.net/howtos/HowTos/",
        "https://acmecse.net/howtos/AuthenticationBetweenCSEs/",
        "https://acmecse.net/howtos/Docker/",
        "https://acmecse.net/howtos/ExperimentalWebSocketBinding/",
        "https://acmecse.net/howtos/ExportResources/",
        "https://acmecse.net/howtos/HowTo-pyenv/",
        "https://acmecse.net/howtos/RaspberryPi/",
        "https://acmecse.net/howtos/EnablingSyntaxHighlighting/",
        "https://acmecse.net/howtos/ServiceProviderRegistration/",
        "https://acmecse.net/howtos/StandAloneWebUI/"
    ],
    "Help": [
        "https://acmecse.net/help/FAQ/",
        "https://acmecse.net/help/NotebookTutorial/",
        "https://acmecse.net/help/oneM2MRecipes/",
        "https://acmecse.net/help/Contributing/"
    ]
    }  
    acmeRepoFiles = {
    "Configuration": [
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/acme/init/acme.ini.default",
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/README.md"
    ],
    "Scripts_and_Automation": [
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/acme/init/init.as",
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/acme/init/startup.as"
    ],
    "Resource_Logic": [
        # These define how ACME handles specific oneM2M resources
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/acme/resources/AE.py",
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/acme/resources/CNT.py",
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/acme/resources/CIN.py",
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/acme/resources/SUB.py",
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/acme/resources/GRP.py"
    ],
    "Validation_and_Types": [
        # This is the "Core Logic" for how ACME validates oneM2M attributes
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/acme/services/Validator.py",
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/acme/etc/Types.py"
    ],
    "Examples_and_Testing": [
        # These show the AI how to actually form a request
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/tests/testAE.py",
        "https://raw.githubusercontent.com/ankraft/ACME-oneM2M-CSE/master/tests/testCNT.py"
    ]
    }
    # Choose what to run:
    # scrap(acmePath, acmeUrls)
    # convertSpecs(specPath)
    # addRepo(repoPath, acmeRepoFiles)